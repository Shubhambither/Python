from openpyxl.workbook import Workbook
from openpyxl import load_workbook
import math
def sri():
    wb = load_workbook('12603996.xlsx')
    ws=wb["Daily Log"]
    # d=ws['A41'].value.strftime("%d-%b-%Y")
    # print("My note on " ,d,"was : "+ws['N41'].value)
    values=[]

    for row in range(7,ws.max_row+1):
        value=ws.cell(row=row,column=2).value
        if value is not None:
            values.append(value)
    # for i in values:
    #     print(i)    
    return math.fsum(values)/len(values)

print(int(sri()))